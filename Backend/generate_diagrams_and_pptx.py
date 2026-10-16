import os
import shutil
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_diagram_image(slide_num, title, filename):
    # 1200x750 High-Resolution Crisp Architectural Canvas
    width, height = 1200, 750
    img = Image.new('RGB', (width, height), color=(15, 23, 42)) # Deep Navy #0f172a
    draw = ImageDraw.Draw(img)

    # Grid background pattern
    for x in range(0, width, 50):
        draw.line([(x, 0), (x, height)], fill=(30, 41, 59), width=1)
    for y in range(0, height, 50):
        draw.line([(0, y), (width, y)], fill=(30, 41, 59), width=1)

    # Outer Border Frame
    draw.rectangle([12, 12, width-12, height-12], outline=(37, 99, 235), width=3)

    # Top Header Banner
    draw.rectangle([20, 20, width-20, 80], fill=(30, 41, 59), outline=(139, 92, 246), width=2)
    draw.text((40, 38), f"FIGURE {slide_num}: {title.upper()}", fill=(255, 255, 255))

    # Helper Drawing Primitives
    def draw_box(x1, y1, x2, y2, text, subtext="", bg_color=(30, 41, 59), border_color=(37, 99, 235), text_color=(255, 255, 255)):
        draw.rectangle([x1, y1, x2, y2], fill=bg_color, outline=border_color, width=2)
        draw.text((x1 + 18, y1 + 14), text, fill=text_color)
        if subtext:
            lines = subtext.split('\n')
            curr_y = y1 + 38
            for line in lines:
                draw.text((x1 + 18, curr_y), line, fill=(148, 163, 184))
                curr_y += 20

    def draw_arrow(x1, y1, x2, y2, label="", color=(59, 130, 246)):
        draw.line([(x1, y1), (x2, y2)], fill=color, width=3)
        if x1 < x2:
            draw.polygon([(x2, y2), (x2-12, y2-7), (x2-12, y2+7)], fill=color)
        elif x1 > x2:
            draw.polygon([(x2, y2), (x2+12, y2-7), (x2+12, y2+7)], fill=color)
        elif y1 < y2:
            draw.polygon([(x2, y2), (x2-7, y2-12), (x2+7, y2-12)], fill=color)
        elif y1 > y2:
            draw.polygon([(x2, y2), (x2-7, y2+12), (x2+7, y2+12)], fill=color)
        
        if label:
            lx = (x1 + x2) // 2
            ly = (y1 + y2) // 2
            draw.rectangle([lx-35, ly-12, lx+35, ly+12], fill=(15, 23, 42), outline=color, width=1)
            draw.text((lx-30, ly-8), label, fill=(203, 213, 225))

    def draw_badge(x, y, text, bg_color=(139, 92, 246), text_color=(255, 255, 255)):
        draw.rectangle([x, y, x+110, y+30], fill=bg_color)
        draw.text((x+12, y+7), text, fill=text_color)

    # -------------------------------------------------------------
    # SLIDE DIAGRAM SPECIFIC GRAPHICS
    # -------------------------------------------------------------
    if slide_num == 1: # Title Overview Platform Hub
        draw_box(450, 310, 750, 430, "EDRP PLATFORM CORE", "Central Decision Hub\nReplay Engine v4.0", (37, 99, 235), (255, 255, 255))
        draw_box(80, 130, 360, 230, "Phase 1: Search Engine", "Multi-Criterion Query Filters\nDebounced Header Dropdown", (30, 41, 59), (59, 130, 246))
        draw_box(840, 130, 1120, 230, "Phase 2: PDF Exports", "ReportLab Vector Engine\nSingle & Executive Reports", (30, 41, 59), (139, 92, 246))
        draw_box(80, 520, 360, 620, "Phase 3: WebSockets", "ConnectionManager Push Feed\nJWT Token Authorization", (30, 41, 59), (13, 148, 136))
        draw_box(840, 520, 1120, 620, "Phase 4: Docker Deploy", "Postgres + FastAPI + Nginx\nHealth Probe Orchestration", (30, 41, 59), (16, 185, 129))

        draw_arrow(360, 180, 470, 310, "Search Queries")
        draw_arrow(730, 310, 840, 180, "PDF Generator")
        draw_arrow(360, 570, 470, 430, "Live Alerts")
        draw_arrow(730, 430, 840, 570, "Containers")

    elif slide_num == 2: # 3-Pillar Executive Objectives
        draw_box(70, 130, 390, 660, "PILLAR 1: KNOWLEDGE BASE", "• Centralized ADR Repository\n• Structured Rationale & Problem\n• Eliminates Tribal Knowledge\n• Decision Category Tagging\n• Searchable Metadata Schema\n• Team Scoped Ownership", (30, 41, 59), (59, 130, 246))
        draw_box(440, 130, 760, 660, "PILLAR 2: TIMELINE REPLAY", "• Step-by-Step Decision History\n• Version Evolution v1 -> vN\n• Comparative Alternative Matrix\n• Risk Level Calculations\n• Interactive Replay Player\n• Change Impact Analysis", (30, 41, 59), (139, 92, 246))
        draw_box(810, 130, 1130, 660, "PILLAR 3: TEAM GOVERNANCE", "• Multi-Stage Reviewer Approval\n• Manager Escalation Engine\n• Team Join Request Controls\n• Manager Full Name Display\n• Dedicated Team Graphs\n• Automated Audit Logging", (30, 41, 59), (16, 185, 129))

    elif slide_num == 3: # 3-Tier Tech Stack
        draw_box(100, 130, 1100, 260, "TIER 1: FRONTEND PRESENTATION LAYER", "React 18 SPA | Vite Build Engine | Vanilla CSS Design Tokens\nCustom SVG Knowledge Graph Canvas | In-App Zoomable Document Viewer | Port 5173", (30, 41, 59), (59, 130, 246))
        draw_arrow(600, 260, 600, 340, "REST / WebSockets (Port 8000)")
        draw_box(100, 340, 1100, 470, "TIER 2: BACKEND APPLICATION LAYER", "FastAPI 0.115+ | Python 3.12 Engine | SQLAlchemy 2.0 ORM | Pydantic v2\nJWT Security | ConnectionManager WebSockets | ReportLab Vector Engine", (30, 41, 59), (139, 92, 246))
        draw_arrow(600, 470, 600, 550, "SQL ORM (Port 5432)")
        draw_box(100, 550, 1100, 680, "TIER 3: DATABASE & INFRASTRUCTURE LAYER", "PostgreSQL 15 Container | Persistent postgres_data Volume | Team Foreign Keys\nDocker Compose Orchestration | Multi-Stage Nginx Alpine Reverse Proxy", (30, 41, 59), (16, 185, 129))

    elif slide_num == 4: # Phase 1 Search Engine
        draw_box(60, 150, 320, 280, "1. USER INPUT MATRIX", "• Search Input Query\n• Category Filter Dropdown\n• Status Pill Selector\n• Custom Date Range", (30, 41, 59), (59, 130, 246))
        draw_arrow(320, 215, 430, 215, "GET /search")
        draw_box(430, 150, 770, 280, "2. FASTAPI SEARCH ENGINE", "• Full-Text Search Handler\n• Title & Problem Matching\n• Rationale & Option Querying\n• Debounced 300ms Input", (30, 41, 59), (139, 92, 246))
        draw_arrow(770, 215, 880, 215, "SQL Query")
        draw_box(880, 150, 1140, 280, "3. POSTGRES DB", "• Indexed Decision Records\n• Foreign Key Scoped\n• Fast Relational Execution", (30, 41, 59), (16, 185, 129))
        draw_box(200, 370, 1000, 650, "FILTERED SEARCH RESULT CARDS (Header.jsx & Search Page)", "✓ Instant drop-down matches under header search bar\n✓ Deep content search across architectural trade-off names\n✓ Sortable by Creation Date (Asc/Desc) and Status Priority", (15, 23, 42), (59, 130, 246))

    elif slide_num == 5: # Phase 2 PDF Generation Engine
        draw_box(80, 140, 560, 340, "EXPORTER 1: SINGLE ADR PDF REPORT", "Endpoint: GET /reports/decisions/{id}/export/pdf\n• Renders Full Title, Author & Category\n• Embedded Problem Statement & Rationale\n• Formatted Alternatives Comparison Table", (30, 41, 59), (59, 130, 246))
        draw_box(640, 140, 1120, 340, "EXPORTER 2: EXECUTIVE SUMMARY PDF", "Endpoint: GET /reports/decisions/export/pdf\n• Multi-Page Executive Summary Document\n• Complete Organization Decision Statistics\n• Team Metrics & Approval Breakdowns", (30, 41, 59), (139, 92, 246))
        draw_arrow(320, 340, 600, 430)
        draw_arrow(880, 340, 600, 430)
        draw_box(300, 430, 900, 650, "REPORTLAB BINARY VECTOR ENGINE", "• High-Resolution PDF Stream Generation\n• Custom Document Styling & Page Numbering\n• Instant 1-Click Browser File Retrieval Download", (30, 41, 59), (16, 185, 129))

    elif slide_num == 6: # Phase 3 WebSockets
        draw_box(60, 140, 340, 290, "1. CLIENT CONNECTION", "• React WS Listener\n• Pass JWT Token Header\n• Automatic Reconnect Loop", (30, 41, 59), (59, 130, 246))
        draw_arrow(340, 215, 440, 215, "ws://.../ws")
        draw_box(440, 140, 760, 290, "2. CONNECTION MANAGER", "• Session Registry\n• Active User Tracking\n• Broadcaster Engine", (30, 41, 59), (139, 92, 246))
        draw_arrow(760, 215, 860, 215, "Push Event")
        draw_box(860, 140, 1140, 290, "3. CLIENT BADGES", "• Animated Toast Popup\n• Unread Badge Counter\n• Zero Refresh Updates", (30, 41, 59), (16, 185, 129))
        draw_box(150, 380, 1050, 660, "EVENT TRIGGER WORKFLOW AUTOMATIONS", "⚡ Decision Status Changes: Approved / Rejected / Escalated alerts author\n⚡ Team Join Requests: Manager notified instantly with inline Action buttons\n⚡ Role Assignments: Employee notified upon manager role change", (15, 23, 42), (13, 148, 136))

    elif slide_num == 7: # PHASE 4 DOCKER DEPLOYMENT MECHANICS (COMPLETE ARCHITECTURE & WORKING SYSTEM PICTURE)
        # Outer Host Machine Boundary
        draw.rectangle([40, 110, 1160, 690], outline=(37, 99, 235), width=3)
        draw.rectangle([40, 110, 1160, 155], fill=(30, 41, 59))
        draw.text((60, 125), "HOST OPERATING SYSTEM — DOCKER ENGINE & DOCKER COMPOSE ORCHESTRATOR", fill=(59, 130, 246))

        # STEP 1: Compose CLI Box
        draw_badge(60, 175, "STEP 1: TRIGGER", (139, 92, 246))
        draw_box(60, 215, 290, 330, "CLI COMMAND", "docker compose up\n  --build -d", (30, 41, 59), (139, 92, 246))
        draw_arrow(290, 270, 350, 270, "1. Launch")

        # Docker Internal Bridge Network Box
        draw.rectangle([350, 175, 1140, 660], outline=(13, 148, 136), width=2)
        draw.text((370, 190), "DOCKER ISOLATED BRIDGE NETWORK (edrp_network)", fill=(13, 148, 136))

        # STEP 2: Database Container
        draw_badge(370, 225, "STEP 2: DATABASE", (16, 185, 129))
        draw_box(370, 265, 610, 620, "edrp_db CONTAINER\n(PostgreSQL 15 Alpine)", "• Image: postgres:15-alpine\n• Port: 5432:5432\n• Health Probe Loop:\n  pg_isready -U edrp_user\n  interval: 5s, timeout: 5s\n• Persistent Volume:\n  postgres_data:/var/lib/...\n• Status: 🟢 HEALTHY", (30, 41, 59), (16, 185, 129))
        draw_arrow(610, 410, 660, 410, "2. Healthy Signal")

        # STEP 3: Backend Container
        draw_badge(660, 225, "STEP 3: BACKEND", (139, 92, 246))
        draw_box(660, 265, 900, 620, "edrp_backend CONTAINER\n(FastAPI / Python 3.12)", "• Image: python:3.12-slim\n• Port: 8000:8000\n• Dependency Guard:\n  depends_on: postgres:\n    condition: service_healthy\n• Server: Uvicorn ASGI\n• Runs DB Migrations\n• Status: 🟢 RUNNING", (30, 41, 59), (139, 92, 246))
        draw_arrow(900, 410, 940, 410, "3. Expose API")

        # STEP 4: Frontend Container
        draw_badge(940, 225, "STEP 4: FRONTEND", (59, 130, 246))
        draw_box(940, 265, 1120, 620, "edrp_frontend\n(Nginx SPA Proxy)", "• Multi-Stage Build:\n  Node 20 -> Nginx\n• Port: 5173:80\n• Serves React SPA\n• Reverse Proxy to\n  edrp_backend:8000\n• Status: 🟢 ACTIVE", (30, 41, 59), (59, 130, 246))

    elif slide_num == 8: # Phase 5 AI Insights
        draw_box(60, 140, 320, 280, "1. DECISION INPUT", "• Rationale & Problem\n• Trade-off Matrix\n• Options & Criteria", (30, 41, 59), (59, 130, 246))
        draw_arrow(320, 210, 430, 210, "JSON Payload")
        draw_box(430, 140, 770, 280, "2. AI ANALYTICS ENGINE", "• Endpoint: GET /ai-summary\n• Risk Score Model\n• Option Choice Ranking", (30, 41, 59), (139, 92, 246))
        draw_arrow(770, 210, 880, 210, "Insights Output")
        draw_box(880, 140, 1140, 280, "3. AI SUMMARY CARD", "• Risk Level Badges\n• Executive Summary\n• Recommended Option", (30, 41, 59), (16, 185, 129))
        draw_box(150, 360, 1050, 650, "DECISION DETAIL INTERACTIVE WIDGET (DecisionDetail.jsx)", "🤖 Executive Trade-off Summary: Synthesizes complex architectural choices into 3 bullet points\n🛡️ Automated Risk Assessment: Assigns Low / Medium / High Risk badges based on complexity\n💡 Optimal Architecture Choice: Highlights top recommended decision alternative with justification", (15, 23, 42), (139, 92, 246))

    elif slide_num == 9: # Team Knowledge Graphs & Document Viewer
        # Team Central Hub Node
        draw.ellipse([510, 240, 690, 420], fill=(139, 92, 246), outline=(255, 255, 255), width=3)
        draw.text((545, 320), "TEAM HUB NODE\n(AI Architecture)", fill=(255, 255, 255))

        # Radiating Decisions
        draw.ellipse([200, 140, 340, 250], fill=(37, 99, 235), outline=(191, 219, 254), width=2)
        draw.text((230, 185), "ADR-101\n(FastAPI)", fill=(255, 255, 255))
        draw_arrow(510, 300, 340, 210, "Team Decision")

        draw.ellipse([860, 140, 1000, 250], fill=(37, 99, 235), outline=(191, 219, 254), width=2)
        draw.text((890, 185), "ADR-102\n(PostgreSQL)", fill=(255, 255, 255))
        draw_arrow(690, 300, 860, 210, "Team Decision")

        # Connected Document Attachments
        draw.ellipse([70, 310, 190, 410], fill=(13, 148, 136), outline=(45, 212, 191), width=2)
        draw.text((95, 350), "📄 spec.pdf", fill=(255, 255, 255))
        draw_arrow(200, 220, 170, 310, "", (45, 212, 191))

        draw.ellipse([1010, 310, 1130, 410], fill=(13, 148, 136), outline=(45, 212, 191), width=2)
        draw.text((1035, 350), "📄 arch.png", fill=(255, 255, 255))
        draw_arrow(1000, 220, 1030, 310, "", (45, 212, 191))

        draw_box(150, 490, 1050, 680, "IN-APP ZOOMABLE DOCUMENT VIEWER MODAL (DocumentViewerModal.jsx)", "🔍 In-App Zoom Controls (+ / - / Reset 100%)\n🔄 90° Canvas Rotation for Blueprints & Diagrams\n📐 Responsive Viewport Scaling & Docked Non-Blocking Inspector Panel\n📥 1-Click File Retrieval & Direct Browser Download", (30, 41, 59), (13, 148, 136))

    elif slide_num == 10: # Quality Assurance & Verification
        draw_box(80, 140, 560, 350, "✅ AUTOMATED PYTEST SUITE PASS", "• Ran 46 out of 46 Pytest Unit Tests\n• 100% Test Pass Rate Across Backend\n• Verifies Auth, Scoping, WebSockets & PDF Exports", (30, 41, 59), (16, 185, 129))
        draw_box(640, 140, 1120, 350, "✅ VITE PRODUCTION ASSET BUILD", "• Command: npx vite build\n• Result: 0 Errors, Compiled in 1.79s\n• Optimized Production Bundle Ready", (30, 41, 59), (16, 185, 129))
        draw_box(80, 410, 560, 660, "✅ POSTGRES MIGRATION & RELATIONAL INTEGRITY", "• Applied team_id column migration\n• Distributed decisions across all 6 Teams\n• Preserved FK cascade & indexing integrity", (30, 41, 59), (59, 130, 246))
        draw_box(640, 410, 1120, 660, "✅ ONE-COMMAND DOCKER DEPLOYMENT", "• Command: docker compose up --build -d\n• Launches PostgreSQL 15, FastAPI & Nginx\n• Complete 3-tier system running seamlessly", (30, 41, 59), (139, 92, 246))

    os.makedirs(os.path.dirname(filename), exist_ok=True)
    img.save(filename)
    print(f"Generated figure diagram: {filename}")

def build_complete_presentation():
    # Step 1: Create 10 Diagram Images
    diagram_dir = "/home/v4u/snap/antigravity/5/.gemini/antigravity/scratch/edrp-milestone4/Backend/static/uploads/figures"
    os.makedirs(diagram_dir, exist_ok=True)

    diagram_meta = [
        (1, "Platform Hub & Milestone 4 Architecture", f"{diagram_dir}/fig1_hub.png"),
        (2, "3-Pillar Executive Objectives", f"{diagram_dir}/fig2_objectives.png"),
        (3, "3-Tier Technology Stack", f"{diagram_dir}/fig3_techstack.png"),
        (4, "Phase 1 Search Engine Pipeline", f"{diagram_dir}/fig4_search.png"),
        (5, "Phase 2 PDF Generation Engine", f"{diagram_dir}/fig5_pdf.png"),
        (6, "Phase 3 WebSockets Live Feed", f"{diagram_dir}/fig6_websockets.png"),
        (7, "Phase 4 Docker System Architecture & Working Mechanics", f"{diagram_dir}/fig7_docker.png"),
        (8, "Phase 5 AI Recommendation Pipeline", f"{diagram_dir}/fig8_ai.png"),
        (9, "Team Knowledge Graph & Document Viewer Topology", f"{diagram_dir}/fig9_teamgraph.png"),
        (10, "Quality Assurance Dashboard & Verification", f"{diagram_dir}/fig10_qa.png")
    ]

    for num, title, filepath in diagram_meta:
        create_diagram_image(num, title, filepath)

    # Step 2: Build PowerPoint Widescreen Presentation (10 Slides with Figures)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    NAVY = RGBColor(15, 23, 42)
    DARK_BLUE = RGBColor(30, 41, 59)
    PRIMARY_BLUE = RGBColor(37, 99, 235)
    LIGHT_BLUE = RGBColor(239, 246, 255)
    PURPLE = RGBColor(139, 92, 246)
    TEAL = RGBColor(13, 148, 136)
    GRAY_TEXT = RGBColor(100, 116, 139)
    WHITE = RGBColor(255, 255, 255)

    def add_header(slide, title_text, category_text="MILESTONE 4 TECHNICAL PRESENTATION"):
        cat_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.3), Inches(12), Inches(0.3))
        tf_cat = cat_box.text_frame
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(10.5)
        p_cat.font.bold = True
        p_cat.font.color.rgb = PRIMARY_BLUE

        title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.55), Inches(12), Inches(0.7))
        tf_title = title_box.text_frame
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = NAVY

    def add_card(slide, left, top, width, height, bg_color=LIGHT_BLUE, border_color=PRIMARY_BLUE):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        if border_color:
            shape.line.color.rgb = border_color
            shape.line.width = Pt(1.5)
        else:
            shape.line.fill.background()
        return shape

    # SLIDE 1: Title Slide
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = NAVY
    bg1.line.fill.background()

    tb1 = s1.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(6.2), Inches(4.5))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "Expert Decision Replay Platform"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p2 = tf1.add_paragraph()
    p2.text = "Milestone 4 Technical Overview & Docker System Mechanics"
    p2.font.size = Pt(18)
    p2.font.color.rgb = PURPLE
    p2.space_before = Pt(14)
    p3 = tf1.add_paragraph()
    p3.text = "Complete 10-Slide Deck with Visual Architecture Diagrams for Every Slide"
    p3.font.size = Pt(13)
    p3.font.color.rgb = GRAY_TEXT
    p3.space_before = Pt(20)

    s1.shapes.add_picture(diagram_meta[0][2], Inches(7.2), Inches(1.5), width=Inches(5.5))

    # SLIDES 2 to 10 Content Data (Left = Text Content, Right = Embedded Visual Figure Diagram)
    slide_contents = [
        # Slide 2
        ("Executive Summary & System Objectives", [
            ("Centralized Knowledge Repository", "Eliminates tribal knowledge by recording structured Architectural Decision Records (ADRs)."),
            ("Multi-Decision Timeline Replay", "Step back and forward through decision evolution history."),
            ("Team Governance & Approval Workflows", "Multi-stage reviewer approvals, escalations, and manager join request controls."),
            ("Milestone 4 Deliverables", "Phases 1-5 complete: Search, PDF Exports, WebSockets, Dockerization, AI Insights & Separate Team Knowledge Graphs.")
        ]),
        # Slide 3
        ("Enterprise Technology Stack & Architecture", [
            ("Frontend Application Tier", "React 18, Vite Build System, Vanilla CSS Token System, Custom SVG Knowledge Graph Canvas, Zoomable In-App Document Viewer."),
            ("Backend Service Tier", "FastAPI 0.115+, Python 3.12 Engine, SQLAlchemy 2.0 ORM, Pydantic v2 Schema Validation, ReportLab Binary PDF Generator."),
            ("Database & Deployment Tier", "PostgreSQL 15 Relational Database, Docker Compose Orchestration, Nginx Alpine Reverse Proxy, 46 Automated Pytest Unit Tests.")
        ]),
        # Slide 4
        ("Phase 1: Multi-Criterion Search & Filtering Engine", [
            ("REST API /decisions/search Endpoint", "Supports Category, Status, Date Range (from_date, to_date), and Keyword search parameters simultaneously."),
            ("Full-Text Deep Querying", "Searches decision titles, problem statements, rationale, and alternative option names/descriptions."),
            ("Global Header Component (Header.jsx)", "Instant debounced drop-down search results available globally across all pages."),
            ("Dynamic Sorting Controls", "Sort by creation date (newest/oldest), title alphabetical, or status priority.")
        ]),
        # Slide 5
        ("Phase 2: PDF Document Generation System", [
            ("Single Decision ADR PDF Export", "GET /reports/decisions/{id}/export/pdf — Renders decision metadata, author details, rationale, and alternatives evaluation matrix."),
            ("Executive Summary PDF Report", "GET /reports/decisions/export/pdf — Multi-page formatted executive report designed for leadership & compliance auditing."),
            ("ReportLab Vector Engine", "Generates high-resolution binary vector PDFs with 1-click browser download integration.")
        ]),
        # Slide 6
        ("Phase 3: Real-Time WebSockets & Push Alerts Feed", [
            ("ConnectionManager (security/websocket_manager.py)", "Manages active user WebSocket sessions with JWT authentication token validation."),
            ("Live Event Broadcasting", "Broadcasts real-time notifications when decisions are approved/rejected, join requests submitted, or team roles assigned."),
            ("Persistent Client Listener (Header.jsx)", "Maintains persistent WebSocket connection with automatic ping/pong keep-alive and reconnection logic."),
            ("Instant Push Alerts", "Displays animated pop-up toast alerts and increments unread notification badge counts without page refreshes.")
        ]),
        # Slide 7 (DOCKER DEPLOYMENT MECHANICS SLIDE)
        ("Phase 4: Docker Containerization & System Mechanics", [
            ("Container Startup & Health Probe", "PostgreSQL container (edrp_db) initializes first; probes pg_isready -U edrp_user -d edrp_db every 5 seconds."),
            ("Backend Dependency Guard", "edrp_backend specifies depends_on: postgres: condition: service_healthy. FastAPI starts ONLY after database is 100% ready."),
            ("Frontend Nginx Production Stage", "edrp_frontend uses multi-stage build (Node 20 build -> Nginx Alpine) serving compiled React SPA on port 5173."),
            ("Network & Volume Persistence", "Containers communicate across internal Docker bridge (postgres:5432). Volume postgres_data preserves database state across restarts."),
            ("Single-Command Execution", "Running 'docker compose up --build -d' compiles assets, creates networks, and launches all 3 containers in background mode.")
        ]),
        # Slide 8
        ("Phase 5: AI Insights & Recommendation Engine", [
            ("GET /decisions/{id}/ai-summary Endpoint", "Analyzes decision rationale, problem statement, and alternatives matrix to return AI-driven insights."),
            ("Executive Technical Summaries", "Generates concise executive summaries highlighting core architectural trade-offs."),
            ("Automated Risk Assessment", "Calculates risk scores and risk level badges (Low, Medium, High)."),
            ("Alternative Option Ranking", "Recommends the optimal architectural choice based on trade-off metrics."),
            ("Interactive UI Component (DecisionDetail.jsx)", "1-click 'Generate AI Insights' widget displaying recommendation cards.")
        ]),
        # Slide 9
        ("Platform Enhancements: Separate Team Graphs & Document Viewer", [
            ("Dedicated Separate Team Knowledge Graphs", "Added team selection pills allowing users to view separate, team-centric Knowledge Graphs for every team in the organization."),
            ("Central Team Hub Nodes", "Renders Team Hub Nodes connecting team decisions and attached supported document nodes."),
            ("In-App Document Viewer Modal (DocumentViewerModal.jsx)", "Opens attached blueprints, PDFs, and diagrams in a clean modal overlay with Zoom (+/-), 90° Rotation, and 1-Click Retrieval Download.")
        ]),
        # Slide 10
        ("Quality Assurance & Project Verification", [
            ("Automated Pytest Suite Health", "Executed complete Backend unit & integration test suite (pytest tests/). Outcome: 46 passed out of 46 tests (100% pass rate)."),
            ("Production Frontend Asset Build", "Compiled production frontend assets using Vite (npx vite build). Outcome: Built in 1.79s with 0 errors."),
            ("PostgreSQL Database Migration", "Applied team_id column migration to PostgreSQL decisions table and verified relational integrity."),
            ("Docker Multi-Container Orchestration", "All 3 containers (PostgreSQL 15, FastAPI, Nginx) verified ready for one-command deployment (docker compose up --build -d).")
        ])
    ]

    for idx, (title, items) in enumerate(slide_contents):
        slide_num = idx + 2
        slide = prs.slides.add_slide(blank_layout)
        add_header(slide, title)

        # Left Column Card for Text Content
        add_card(slide, Inches(0.6), Inches(1.4), Inches(6.5), Inches(5.6), WHITE, PRIMARY_BLUE)
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(6.1), Inches(5.3))
        tf = tb.text_frame
        tf.word_wrap = True

        for item_idx, item in enumerate(items):
            if isinstance(item, tuple):
                t_head, t_desc = item
                p = tf.paragraphs[0] if item_idx == 0 else tf.add_paragraph()
                p.text = f"• {t_head}: "
                p.font.size = Pt(13)
                p.font.bold = True
                p.font.color.rgb = PRIMARY_BLUE
                if item_idx > 0:
                    p.space_before = Pt(10)

                # Append description
                run = p.add_run()
                run.text = t_desc
                run.font.bold = False
                run.font.color.rgb = DARK_BLUE
            else:
                p = tf.paragraphs[0] if item_idx == 0 else tf.add_paragraph()
                p.text = f"• {item}"
                p.font.size = Pt(12.5)
                p.font.color.rgb = DARK_BLUE
                if item_idx > 0:
                    p.space_before = Pt(10)

        # Right Column: Embed Supported Visual Figure Diagram Image
        fig_path = diagram_meta[idx + 1][2]
        slide.shapes.add_picture(fig_path, Inches(7.3), Inches(1.4), width=Inches(5.4))

    # Save to all required paths
    output_path = "/home/v4u/snap/antigravity/5/.gemini/antigravity/scratch/edrp-milestone4/Backend/static/uploads/EDRP_Milestone4_Presentation.pptx"
    root_path = "/home/v4u/snap/antigravity/5/.gemini/antigravity/scratch/edrp-milestone4/EDRP_Milestone4_Presentation.pptx"
    public_path = "/home/v4u/snap/antigravity/5/.gemini/antigravity/scratch/edrp-milestone4/frontend/public/EDRP_Milestone4_Presentation.pptx"

    prs.save(output_path)
    shutil.copyfile(output_path, root_path)
    os.makedirs(os.path.dirname(public_path), exist_ok=True)
    shutil.copyfile(output_path, public_path)

    print("=========================================================")
    print("SUCCESSFULLY GENERATED 10-SLIDE PRESENTATION WITH FIGURES!")
    print(f"PPTX saved to: {output_path}")
    print(f"PPTX copied to: {root_path}")
    print(f"PPTX copied to: {public_path}")
    print("=========================================================")

if __name__ == "__main__":
    build_complete_presentation()
