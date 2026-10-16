import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # Set slide dimensions to widescreen 16:9 (13.33 x 7.5 inches)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_slide_layout = prs.slide_layouts[6]

    # Color Palette
    NAVY = RGBColor(15, 23, 42)      # #0f172a
    DARK_BLUE = RGBColor(30, 41, 59)  # #1e293b
    PRIMARY_BLUE = RGBColor(37, 99, 235) # #2563eb
    LIGHT_BLUE = RGBColor(239, 246, 255) # #eff6ff
    PURPLE = RGBColor(139, 92, 246)  # #8b5cf6
    TEAL = RGBColor(13, 148, 136)    # #0d9488
    GRAY_TEXT = RGBColor(100, 116, 139) # #64748b
    WHITE = RGBColor(255, 255, 255)
    LIGHT_BG = RGBColor(248, 250, 252) # #f8fafc

    def add_header(slide, title_text, category_text="MILESTONE 4 TECHNICAL PRESENTATION"):
        # Category Tracker
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11), Inches(0.4))
        tf_cat = cat_box.text_frame
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = PRIMARY_BLUE

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11), Inches(0.8))
        tf_title = title_box.text_frame
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(26)
        p_title.font.bold = True
        p_title.font.color.rgb = NAVY

    def add_card(slide, left, top, width, height, bg_color=LIGHT_BG, border_color=PRIMARY_BLUE):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        if border_color:
            shape.line.color.rgb = border_color
            shape.line.width = Pt(1.5)
        else:
            shape.line.fill.background()
        return shape

    # ==========================================
    # SLIDE 1: Title Slide (Dark Theme)
    # ==========================================
    slide1 = prs.slides.add_slide(blank_slide_layout)
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = NAVY
    bg1.line.fill.background()

    # Title Card Accent
    accent = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(2.2), Inches(0.15), Inches(3.2))
    accent.fill.solid()
    accent.fill.fore_color.rgb = PRIMARY_BLUE
    accent.line.fill.background()

    tb1 = slide1.shapes.add_textbox(Inches(1.2), Inches(2.1), Inches(11), Inches(3.5))
    tf1 = tb1.text_frame
    p1 = tf1.paragraphs[0]
    p1.text = "Expert Decision Replay Platform"
    p1.font.size = Pt(40)
    p1.font.bold = True
    p1.font.color.rgb = WHITE

    p2 = tf1.add_paragraph()
    p2.text = "Milestone 4 Comprehensive Technical Overview & Architecture Demonstration"
    p2.font.size = Pt(22)
    p2.font.color.rgb = PURPLE
    p2.space_before = Pt(14)

    p3 = tf1.add_paragraph()
    p3.text = "Includes Phases 1–5: Search, PDF Exports, WebSockets, Docker Deployment Mechanics, AI Insights & Team Graphs"
    p3.font.size = Pt(14)
    p3.font.color.rgb = GRAY_TEXT
    p3.space_before = Pt(24)

    # ==========================================
    # SLIDE 2: Executive Summary & Project Goal
    # ==========================================
    slide2 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide2, "Executive Summary & System Objectives")

    # Left Box - Problem & Mission
    add_card(slide2, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0), WHITE, PRIMARY_BLUE)
    tb2_left = slide2.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.2), Inches(4.6))
    tf2_l = tb2_left.text_frame
    tf2_l.word_wrap = True
    
    p = tf2_l.paragraphs[0]
    p.text = "🎯 Core Objective & Mission"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = NAVY

    items_left = [
      ("Centralized Knowledge Repository", "Eliminates tribal knowledge by recording structured Architectural Decision Records (ADRs)."),
      ("Multi-Decision Timeline Replay", "Allows engineering teams to step back and forward through version evolution history."),
      ("Team Governance & Approval Workflows", "Multi-stage reviewer approvals, escalations, and manager join request controls.")
    ]
    for title, desc in items_left:
        p = tf2_l.add_paragraph()
        p.text = f"• {title}: {desc}"
        p.font.size = Pt(13)
        p.font.color.rgb = DARK_BLUE
        p.space_before = Pt(14)

    # Right Box - Milestone 4 Capabilities
    add_card(slide2, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0), LIGHT_BLUE, PURPLE)
    tb2_right = slide2.shapes.add_textbox(Inches(7.0), Inches(2.0), Inches(5.3), Inches(4.6))
    tf2_r = tb2_right.text_frame
    tf2_r.word_wrap = True

    p = tf2_r.paragraphs[0]
    p.text = "🚀 Milestone 4 Production Deliverables"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = NAVY

    items_right = [
      "Phase 1: Multi-Criterion Search & Filtering Engine",
      "Phase 2: ReportLab Binary PDF Export System",
      "Phase 3: Real-Time WebSockets Live Feed & Push Alerts",
      "Phase 4: Docker Containerization & Multi-Container Deployment",
      "Phase 5: AI Architectural Summary & Recommendations Engine",
      "Enhancement: Separate Team Knowledge Graphs & Document Viewer Modal"
    ]
    for item in items_right:
        p = tf2_r.add_paragraph()
        p.text = f"✅ {item}"
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = DARK_BLUE
        p.space_before = Pt(12)

    # ==========================================
    # SLIDE 3: Technical Stack & Architecture
    # ==========================================
    slide3 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide3, "Enterprise Technology Stack & Architecture")

    stacks = [
        ("Frontend Application", PRIMARY_BLUE, ["React 18 & Vite Build System", "Vanilla CSS Token System", "React Router 6 Navigation", "Custom SVG Knowledge Graph Canvas", "Zoomable In-App Document Viewer"]),
        ("Backend FastAPI Service", PURPLE, ["FastAPI 0.115+ REST & WebSocket API", "Python 3.12 Engine", "SQLAlchemy 2.0 ORM Engine", "Pydantic v2 Schema Validation", "ReportLab Binary PDF Generator"]),
        ("Database & Infrastructure", TEAL, ["PostgreSQL 15 Relational Database", "Foreign Key Integrity & Cascades", "Docker & Docker Compose Orchestration", "Nginx Reverse Proxy & Asset Server", "46 Automated Pytest Unit Tests"])
    ]

    for idx, (title, color, bullets) in enumerate(stacks):
        left_pos = Inches(0.8 + idx * 4.0)
        add_card(slide3, left_pos, Inches(1.8), Inches(3.7), Inches(5.0), WHITE, color)
        
        tb = slide3.shapes.add_textbox(left_pos + Inches(0.2), Inches(2.0), Inches(3.3), Inches(4.6))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(17)
        p.font.bold = True
        p.font.color.rgb = color
        
        for b in bullets:
            p = tf.add_paragraph()
            p.text = f"▪ {b}"
            p.font.size = Pt(12.5)
            p.font.color.rgb = DARK_BLUE
            p.space_before = Pt(12)

    # ==========================================
    # SLIDE 4: Milestone 4 Phase 1 — Search Engine
    # ==========================================
    slide4 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide4, "Phase 1: Multi-Criterion Search & Filtering Engine")

    add_card(slide4, Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0), WHITE, PRIMARY_BLUE)
    tb4 = slide4.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(11.1), Inches(4.6))
    tf4 = tb4.text_frame
    tf4.word_wrap = True

    p = tf4.paragraphs[0]
    p.text = "🔍 REST API `/decisions/search` & Live Header Bar Integration"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_BLUE

    p_body = [
      "Multi-Criterion Search Parameters: Supports Category, Status, Date Range (from_date, to_date), and Keyword search simultaneously.",
      "Full-Text Deep Querying: Searches decision titles, problem statements, rationale, and alternative option names/descriptions.",
      "Live Header Bar Component (Header.jsx): Instant debounced drop-down search results available globally across all application pages.",
      "Dynamic Sorting Options: Sort by creation date (newest/oldest), title alphabetical, or status priority."
    ]
    for b in p_body:
        p = tf4.add_paragraph()
        p.text = f"• {b}"
        p.font.size = Pt(14)
        p.font.color.rgb = DARK_BLUE
        p.space_before = Pt(16)

    # ==========================================
    # SLIDE 5: Milestone 4 Phase 2 — PDF Export Generator
    # ==========================================
    slide5 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide5, "Phase 2: PDF Document Generation System")

    # Card 1: Single Decision PDF
    add_card(slide5, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0), WHITE, PRIMARY_BLUE)
    tb5_l = slide5.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.2), Inches(4.6))
    tf5_l = tb5_l.text_frame
    tf5_l.word_wrap = True
    p = tf5_l.paragraphs[0]
    p.text = "📄 Single Decision ADR PDF Export"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_BLUE
    for b in [
      "Endpoint: GET /reports/decisions/{id}/export/pdf",
      "Engine: ReportLab binary document generator",
      "Includes decision metadata, author details, status badge, problem statement, rationale, and alternatives evaluation table.",
      "Instant 1-click download from Decision Detail page."
    ]:
        p = tf5_l.add_paragraph()
        p.text = f"▪ {b}"
        p.font.size = Pt(13)
        p.font.color.rgb = DARK_BLUE
        p.space_before = Pt(14)

    # Card 2: Executive Summary PDF
    add_card(slide5, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0), WHITE, PURPLE)
    tb5_r = slide5.shapes.add_textbox(Inches(7.0), Inches(2.0), Inches(5.3), Inches(4.6))
    tf5_r = tb5_r.text_frame
    tf5_r.word_wrap = True
    p = tf5_r.paragraphs[0]
    p.text = "📊 Executive Summary PDF Report"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = PURPLE
    for b in [
      "Endpoint: GET /reports/decisions/export/pdf",
      "Multi-page formatted executive report designed for leadership & compliance auditing.",
      "Renders organization-wide decision statistics, team breakdowns, and approved technical specifications.",
      "Integrated into Reports dashboard page."
    ]:
        p = tf5_r.add_paragraph()
        p.text = f"▪ {b}"
        p.font.size = Pt(13)
        p.font.color.rgb = DARK_BLUE
        p.space_before = Pt(14)

    # ==========================================
    # SLIDE 6: Milestone 4 Phase 3 — Real-Time WebSockets
    # ==========================================
    slide6 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide6, "Phase 3: Real-Time WebSockets & Notifications Feed")

    add_card(slide6, Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0), WHITE, TEAL)
    tb6 = slide6.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(11.1), Inches(4.6))
    tf6 = tb6.text_frame
    tf6.word_wrap = True

    p = tf6.paragraphs[0]
    p.text = "⚡ Authenticated WebSocket Connection Manager"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = TEAL

    p_body6 = [
      "ConnectionManager (security/websocket_manager.py): Manages active user WebSocket sessions with JWT authentication token validation.",
      "Live Event Broadcasting: Broadcasts real-time notifications when decisions are approved/rejected, join requests submitted, or team roles assigned.",
      "Client Header Listener (Header.jsx): Maintains persistent WebSocket connection with automatic ping/pong keep-alive and reconnection logic.",
      "Instant Live UI Push: Displays animated pop-up toast alerts and dynamically increments unread notification badge counts without requiring page refreshes."
    ]
    for b in p_body6:
        p = tf6.add_paragraph()
        p.text = f"• {b}"
        p.font.size = Pt(13.5)
        p.font.color.rgb = DARK_BLUE
        p.space_before = Pt(16)

    # ==========================================
    # SLIDE 7: Milestone 4 Phase 4 — Docker Containerization
    # ==========================================
    slide7 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide7, "Phase 4: Docker Containerization Architecture")

    containers = [
      ("Database Container (edrp_db)", "postgres:15-alpine image, persistent postgres_data volume, pg_isready health checks."),
      ("Backend Container (edrp_backend)", "Python 3.12 slim image, FastAPI service, dependencies pre-compiled, Uvicorn server on port 8000."),
      ("Frontend Container (edrp_frontend)", "Multi-stage Dockerfile (Node 20 build stage -> Nginx Alpine production stage on port 5173).")
    ]
    for idx, (title, desc) in enumerate(containers):
        top_pos = Inches(1.8 + idx * 1.6)
        add_card(slide7, Inches(0.8), top_pos, Inches(11.7), Inches(1.3), WHITE, PRIMARY_BLUE)
        tb = slide7.shapes.add_textbox(Inches(1.1), top_pos + Inches(0.1), Inches(11.1), Inches(1.1))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f"🐳 {title}"
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = PRIMARY_BLUE
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(13)
        p2.font.color.rgb = DARK_BLUE
        p2.space_before = Pt(4)

    # ==========================================
    # SLIDE 8: How Docker Deployment Works (Step-by-Step Mechanics)
    # ==========================================
    slide8_dep = prs.slides.add_slide(blank_slide_layout)
    add_header(slide8_dep, "How the Docker Deployment Works Step-by-Step")

    add_card(slide8_dep, Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0), LIGHT_BLUE, PRIMARY_BLUE)
    tb8_dep = slide8_dep.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(11.1), Inches(4.6))
    tf8_dep = tb8_dep.text_frame
    tf8_dep.word_wrap = True

    p = tf8_dep.paragraphs[0]
    p.text = "⚙️ Container Startup Sequence, Health Checks & Request Routing"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_BLUE

    steps = [
      ("1. Database Startup & Health Probe", "PostgreSQL container (edrp_db) initializes first. Probes every 5s using pg_isready -U edrp_user -d edrp_db."),
      ("2. Backend Dependency Guard & Launch", "edrp_backend uses depends_on: postgres: condition: service_healthy. FastAPI starts ONLY when PostgreSQL is 100% ready."),
      ("3. Multi-Stage Frontend Build & Nginx Serving", "edrp_frontend builds React SPA with Node 20, copies static bundle to Nginx Alpine, and serves at http://localhost:5173."),
      ("4. Internal Network & Persistent Volumes", "Containers communicate over internal Docker network (postgres:5432). The postgres_data volume preserves database state across restarts."),
      ("5. Single-Command Deployment", "Executing 'docker compose up --build -d' compiles assets, builds containers, and launches the entire platform in background mode.")
    ]
    for title, desc in steps:
        p = tf8_dep.add_paragraph()
        p.text = f"• {title}: {desc}"
        p.font.size = Pt(12.5)
        p.font.color.rgb = DARK_BLUE
        p.space_before = Pt(10)

    # ==========================================
    # SLIDE 9: Milestone 4 Phase 5 — AI Insights
    # ==========================================
    slide8 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide8, "Phase 5: AI Insights & Recommendation Engine")

    add_card(slide8, Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0), WHITE, PURPLE)
    tb8 = slide8.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(11.1), Inches(4.6))
    tf8 = tb8.text_frame
    tf8.word_wrap = True

    p = tf8.paragraphs[0]
    p.text = "🤖 AI Architectural Summary & Risk Assessment Engine"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = PURPLE

    p_body8 = [
      "Endpoint GET /decisions/{id}/ai-summary: Analyzes decision rationale, problem statement, and alternatives matrix to return AI-driven insights.",
      "Executive Technical Summaries: Generates concise executive summaries highlighting core architectural trade-offs.",
      "Automated Risk Assessment: Calculates risk scores and risk level badges (Low, Medium, High).",
      "Alternative Option Ranking: Recommends the optimal architectural choice based on trade-off metrics.",
      "Interactive UI Component (DecisionDetail.jsx): 1-click 'Generate AI Insights' widget displaying recommendation cards."
    ]
    for b in p_body8:
        p = tf8.add_paragraph()
        p.text = f"• {b}"
        p.font.size = Pt(13.5)
        p.font.color.rgb = DARK_BLUE
        p.space_before = Pt(14)

    # ==========================================
    # SLIDE 10: Platform Enhancements — Team Graphs & Document Viewer
    # ==========================================
    slide9 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide9, "Platform Enhancements: Separate Team Knowledge Graphs & Document Viewer")

    # Card 1: Separate Team Graphs
    add_card(slide9, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0), WHITE, PRIMARY_BLUE)
    tb9_l = slide9.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.2), Inches(4.6))
    tf9_l = tb9_l.text_frame
    tf9_l.word_wrap = True
    p = tf9_l.paragraphs[0]
    p.text = "🕸️ Separate Team Knowledge Graphs"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_BLUE
    for b in [
      "Dedicated Team Selection Pills: Allows users to view separate, team-centric Knowledge Graphs for every team in the organization.",
      "Central Team Hub Nodes: Renders Team Hub Nodes connecting team decisions and attached supported document nodes.",
      "Explicit Database Team Scoping: Decisions explicitly associated with teams in PostgreSQL database.",
      "Layout Options: Split Grid View, Full Width Graph, and 1-Click Fullscreen Viewport."
    ]:
        p = tf9_l.add_paragraph()
        p.text = f"▪ {b}"
        p.font.size = Pt(12.5)
        p.font.color.rgb = DARK_BLUE
        p.space_before = Pt(12)

    # Card 2: Document Viewer Modal
    add_card(slide9, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0), WHITE, TEAL)
    tb9_r = slide9.shapes.add_textbox(Inches(7.0), Inches(2.0), Inches(5.3), Inches(4.6))
    tf9_r = tb9_r.text_frame
    tf9_r.word_wrap = True
    p = tf9_r.paragraphs[0]
    p.text = "🔍 Zoomable Document Viewer Modal"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = TEAL
    for b in [
      "In-App Document Previewer (DocumentViewerModal.jsx): Opens attached blueprints, PDFs, and diagrams in a clean modal overlay.",
      "Viewport Controls: Zoom In (+), Zoom Out (-), Reset (100%), and 90° Rotation.",
      "Responsive Fit: Automatically scales documents to fit comfortably within the viewport.",
      "100% Retrievable: Includes prominent '📥 Download File' button retrieving original file."
    ]:
        p = tf9_r.add_paragraph()
        p.text = f"▪ {b}"
        p.font.size = Pt(12.5)
        p.font.color.rgb = DARK_BLUE
        p.space_before = Pt(12)

    # ==========================================
    # SLIDE 11: Verification & Conclusion (Dark Theme)
    # ==========================================
    slide10 = prs.slides.add_slide(blank_slide_layout)
    bg10 = slide10.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg10.fill.solid()
    bg10.fill.fore_color.rgb = NAVY
    bg10.line.fill.background()

    tb10 = slide10.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.7), Inches(6.0))
    tf10 = tb10.text_frame
    tf10.word_wrap = True

    p = tf10.paragraphs[0]
    p.text = "Quality Assurance & Project Verification"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = WHITE

    verifications = [
      ("Automated Pytest Suite Health", "Executed complete Backend unit & integration test suite (pytest tests/). Outcome: 46 passed out of 46 tests (100% pass rate)."),
      ("Production Frontend Asset Build", "Compiled production frontend assets using Vite (npx vite build). Outcome: Built in 1.79s with 0 errors."),
      ("PostgreSQL Database Migration", "Applied team_id column migration to PostgreSQL decisions table and verified relational integrity."),
      ("Docker Multi-Container Orchestration", "All 3 containers (PostgreSQL 15, FastAPI, Nginx) verified ready for one-command deployment (docker compose up --build).")
    ]
    for title, desc in verifications:
        p = tf10.add_paragraph()
        p.text = f"✅ {title}: {desc}"
        p.font.size = Pt(14)
        p.font.color.rgb = LIGHT_BLUE
        p.space_before = Pt(18)

    p_end = tf10.add_paragraph()
    p_end.text = "🎉 EDRP Milestone 4 is 100% Complete, Fully Tested, and Ready for Deployment!"
    p_end.font.size = Pt(18)
    p_end.font.bold = True
    p_end.font.color.rgb = PURPLE
    p_end.space_before = Pt(28)

    output_path = "/home/v4u/snap/antigravity/5/.gemini/antigravity/scratch/edrp-milestone4/Backend/static/uploads/EDRP_Milestone4_Presentation.pptx"
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    prs.save(output_path)
    print(f"Presentation created successfully at: {output_path}")

    # Copy to main project root and frontend public folder
    root_path = "/home/v4u/snap/antigravity/5/.gemini/antigravity/scratch/edrp-milestone4/EDRP_Milestone4_Presentation.pptx"
    public_path = "/home/v4u/snap/antigravity/5/.gemini/antigravity/scratch/edrp-milestone4/frontend/public/EDRP_Milestone4_Presentation.pptx"
    import shutil
    shutil.copyfile(output_path, root_path)
    os.makedirs(os.path.dirname(public_path), exist_ok=True)
    shutil.copyfile(output_path, public_path)

if __name__ == "__main__":
    create_presentation()
